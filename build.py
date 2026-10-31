"""Floating の組み立て：src/ の部品と data/ を1つのHTML（docs/index.html）にまとめる。

  src/shell.html  … HTML・CSS・画面の部品（__SCRIPT__ の位置に描画のコードが入る）
  src/js/NN-*.js  … 描画のコード。ファイル名の番号順につなげる（1つのモジュールとして動く）
  data/lake.json, data/dem.json, data/depth.json, data/isles.json, data/rivers.json, data/towns.json, data/near.json, data/landuse.json, data/rail.json, data/buoys.xlsx … 湖・地形・深さ・線路・ブイのデータ

読み込む量を減らすため（見た目は変えない）：
  ・大きな2進のデータ（base64 の文字列）は、並べ方を変えてから圧縮して埋め込み、作品の中で元の文字列に戻す（pack と src/js/00a-data.js）
  ・コードのコメントと行頭の字下げを取り除く（tools/strip.mjs。TypeScript を使う：npm install -g typescript@6.0.2）
"""
import base64, pathlib, json, re, zlib
from openpyxl import load_workbook

r = pathlib.Path(__file__).parent


LAT_RANGE, LON_RANGE = (34.85, 35.6), (135.8, 136.35)   # 琵琶湖とその周り
YT = re.compile(r'^https?://(?:www\.|m\.)?(?:youtube\.com/(?:watch\?v=|shorts/|live/)|youtu\.be/)([\w-]{11})')


def buoys():
    """ブイ一覧を読み、入力を確かめる。作品に出せない行は飛ばし、気になる点は知らせる。"""
    ws = load_workbook(r / 'data/buoys.xlsx', data_only=True)['ブイ一覧']
    out, notes = [], []
    seen_pos, seen_vid = {}, {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        no, name, form, lat, lon, url, show = row[:7]
        if lat is None and lon is None and not url:
            continue
        tag = f'No.{no}' + (f'（{name}）' if name else '')
        if show == '×':
            continue
        if not isinstance(lat, (int, float)) or not isinstance(lon, (int, float)):
            notes.append(f'{tag}：北緯・東経が数値でないため飛ばしました'); continue
        if LON_RANGE[0] <= lat <= LON_RANGE[1] and LAT_RANGE[0] <= lon <= LAT_RANGE[1]:
            notes.append(f'{tag}：北緯と東経が逆のようです（入れ替えて使いました）'); lat, lon = lon, lat
        if not (LAT_RANGE[0] <= lat <= LAT_RANGE[1] and LON_RANGE[0] <= lon <= LON_RANGE[1]):
            notes.append(f'{tag}：琵琶湖の範囲外の座標（{lat}, {lon}）のため飛ばしました'); continue
        if not url:
            notes.append(f'{tag}：リンクがないため飛ばしました'); continue
        url = str(url).strip()
        m = YT.match(url)
        if not m:
            notes.append(f'{tag}：YouTubeのリンクではないようです（そのまま使いました）：{url}')
        else:
            vid = m.group(1)
            if vid in seen_vid:
                notes.append(f'{tag}：{seen_vid[vid]} と同じ動画です')
            seen_vid[vid] = tag
        key = (round(lat, 4), round(lon, 4))   # 約10mの精度で同じ場所
        if key in seen_pos:
            notes.append(f'{tag}：{seen_pos[key]} とほぼ同じ場所です（約10m以内）')
        seen_pos[key] = tag
        if not name:
            notes.append(f'{tag}：集落名が空です')
        out.append({'no': no, 'name': name or '', 'form': form or '', 'lat': lat, 'lon': lon, 'url': url})
    for n in notes:
        print('ブイ一覧：', n)
    return out


# ---- データを詰める（src/js/00a-data.js で元に戻す）----
def _delta(b):
    o, p = bytearray(len(b)), 0
    for i, v in enumerate(b):
        o[i] = (v - p) & 255
        p = v
    return bytes(o)


def _planes(b, k, d=False):
    n = len(b) // k
    return b''.join(_delta(b[j::k][:n]) if d else bytes(b[j::k][:n]) for j in range(k)) + b[n * k:]


def _pred2d(b, w):
    o = bytearray(len(b))
    for i in range(len(b)):
        a = b[i - 1] if i % w else 0
        c = b[i - w] if i >= w else 0
        d = b[i - w - 1] if i >= w and i % w else 0
        o[i] = (b[i] - (a + c - d)) & 255
    return bytes(o)


def pack(s, t, w=0):
    """base64 の文字列 s を、並べ方 t で並べ替えて圧縮した形にする（どの並べ方が効くかは、データごとに試して決めた）"""
    b = base64.b64decode(s)
    x = b if t == 'raw' else _delta(b) if t == 'delta' else _pred2d(b, w) if t == '2d' else _planes(b, int(t[2:] if t[1] == 'd' else t[1:]), t[1] == 'd')
    c = zlib.compressobj(9, zlib.DEFLATED, -15)
    out = {'$p': base64.b64encode(c.compress(x) + c.flush()).decode(), 't': t}
    if t == '2d':
        out['w'] = w
    return out


def data():
    """作品のコードが DATA として読むデータ（00a-data.js）"""
    L = lambda n: json.loads((r / f'data/{n}.json').read_text())
    dem, depth, near, towns, lu, isles = L('dem'), L('depth'), L('near'), L('towns'), L('landuse'), L('isles')
    dem['d'] = pack(dem['d'], '2d', dem['w'])
    depth['d'] = pack(depth['d'], '2d', depth['w'])
    near['d'], near['idx'] = pack(near['d'], 'raw'), pack(near['idx'], 'raw')
    towns['b'], towns['r'] = pack(towns['b'], 'p8'), pack(towns['r'], 'p2')
    lu['d'] = pack(lu['d'], 'pd5')
    for q in isles:
        q['d'] = pack(q['d'], 'delta')
    return {'dem': dem, 'depth': depth, 'near': near, 'towns': towns, 'landuse': lu, 'isles': isles}


def strip(script):
    """コメントと行頭の字下げを取り除く（tools/strip.mjs）。動かない環境では止める（公開ファイルが環境で変わらないように）"""
    import subprocess
    res = subprocess.run(['node', str(r / 'tools/strip.mjs')], input=script, capture_output=True, text=True)
    if res.returncode or not res.stdout.strip():
        print(res.stderr[-600:])
        raise SystemExit('コードを軽くする手順（tools/strip.mjs）が動きません。npm install -g typescript@6.0.2 を入れてください')
    return res.stdout


def build():
    import sys
    sys.path.insert(0, str(r / 'tools'))
    import deps
    errs = deps.write()
    if errs:
        for e in errs:
            print('エラー：', e)
        raise SystemExit('部品の名前が重なっています（src/DEPENDENCIES.md 参照）')
    parts = sorted((r / 'src/js').glob('*.js'))
    script = strip('\n'.join(p.read_text() for p in parts))
    html = (r / 'src/shell.html').read_text().replace('__SCRIPT__', script)
    B = buoys()
    html = (html.replace('__LAKE__', (r / 'data/lake.json').read_text())
                .replace('__RAIL__', (r / 'data/rail.json').read_text())
                .replace('__RIVERS__', (r / 'data/rivers.json').read_text())
                .replace('__DATA__', json.dumps(data(), separators=(',', ':')))
                .replace('__BUOYS__', json.dumps(B, ensure_ascii=False)))
    # 文法の確認：node があれば、組み立てたコードを読み込めるか確かめる（壊れたまま公開しないため）
    import subprocess, shutil, tempfile
    if shutil.which('node'):
        with tempfile.NamedTemporaryFile('w', suffix='.mjs', delete=False) as f:
            f.write(script.replace('__LAKE__', '{}').replace('__RAIL__', '{}').replace('__RIVERS__', '[]').replace('__DATA__', '{}').replace('__BUOYS__', '[]'))
        res = subprocess.run(['node', '--check', f.name], capture_output=True, text=True)
        if res.returncode:
            print(res.stderr.split('SyntaxError')[0][-400:] + 'SyntaxError' + res.stderr.split('SyntaxError')[-1][:200])
            raise SystemExit('コードに文法の誤りがあるため、組み立てを中止しました')
    (r / 'docs/index.html').write_text(html)
    print(f'built docs/index.html  部品 {len(parts)} 個 / ブイ {len(B)} 個')


if __name__ == '__main__':
    build()
